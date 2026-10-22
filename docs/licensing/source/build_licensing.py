from pathlib import Path
import json,re
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT

R=Path(__file__).resolve().parent;OUT=R/'output';OUT.mkdir(exist_ok=True)
DATA=R/'source' if (R/'source/FACTS.json').exists() else R
F=json.loads((DATA/'FACTS.json').read_text())
C=json.loads((DATA/'LICENSE_CLAUSES.json').read_text())
T=json.loads((DATA/'PAIRED_TEXT.json').read_text());K=F['key']
d=Document();s=d.sections[0]
s.page_width=Inches(8.5);s.page_height=Inches(11)
s.top_margin=Inches(.55);s.bottom_margin=Inches(.52)
s.left_margin=Inches(.65);s.right_margin=Inches(.65)
s.header_distance=Inches(.22);s.footer_distance=Inches(.22)
for x in list(d.styles.element.iter(qn('w:pBdr'))):x.getparent().remove(x)
for name in ['Normal','Title','Subtitle','Heading 1']:
    st=d.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
    st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Times New Roman')
n=d.styles['Normal'];n.font.size=Pt(11);n.paragraph_format.line_spacing=1.01
n.paragraph_format.space_after=Pt(2);n.paragraph_format.widow_control=True
d.styles['Title'].font.size=Pt(20);d.styles['Title'].font.bold=True
d.styles['Title'].paragraph_format.space_after=Pt(5)
d.styles['Subtitle'].font.size=Pt(10);d.styles['Subtitle'].paragraph_format.space_after=Pt(5)
st=d.styles['Heading 1'];st.font.size=Pt(11);st.font.bold=True
st.paragraph_format.space_before=Pt(5);st.paragraph_format.space_after=Pt(2)
st.paragraph_format.keep_with_next=True
hp=s.header.paragraphs[0];hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
hp.add_run('GRANDINIU LD  |  RIMAS KARPAVIČIUS  |  1.1  |  2026 10 22').font.size=Pt(8)
fp=s.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
fp.add_run('EN / LT  ·  ').font.size=Pt(8)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');fp._p.append(field)
d.core_properties.title='Grandiniu LD licencija ir versijų patikra'
d.core_properties.author='Rimas Karpavičius';d.core_properties.language='en-US;lt-LT'
d.core_properties.subject='Equivalent English and Lithuanian texts version 1.1'
coverage={'en':[],'lt':[]}
lang='en'
def p(text,style=None,key=None):
    pp=d.add_paragraph(text,style)
    pr=pp._p.get_or_add_pPr();runpr=OxmlElement('w:rPr');lg=OxmlElement('w:lang');lg.set(qn('w:val'),'en-US' if lang=='en' else 'lt-LT');runpr.append(lg);pr.append(runpr)
    if key:coverage[lang].append(key)
    return pp
def h(text,key=None):return p(text,'Heading 1',key)
def unit(key,heading=False):return h(T[key][lang],key) if heading else p(T[key][lang],key=key)
def mono(text):
    pp=p(text)
    for rr in pp.runs:rr.font.name='DejaVu Sans Mono';rr.font.size=Pt(9)
    return pp
def link(label,url):
    pp=p('');hy=OxmlElement('w:hyperlink');hy.set(qn('r:id'),pp.part.relate_to(url,RT.HYPERLINK,is_external=True))
    rr=OxmlElement('w:r');tt=OxmlElement('w:t');tt.text=label;rr.append(tt);hy.append(rr);pp._p.append(hy)
def page(title,sub):
    if d.paragraphs:d.add_page_break()
    p(title,'Title');p(sub,'Subtitle')
for lang in ['en','lt']:
    page('Grandiniu LD proprietary license' if lang=='en' else 'Grandiniu LD nuosavybinė licencija',('Version 1.1  |  Notice date ' if lang=='en' else '1.1 versija  |  Pranešimo data ')+F['date'])
    p(('Copyright (c) 2026 ' if lang=='en' else 'Autorių teisės (c) 2026 ')+F['owner']+'.')
    p('All rights reserved in the original protectable contributions.' if lang=='en' else 'Visos teisės į originalų teisiškai saugomą indėlį saugomos.')
    p(('Permission contact: ' if lang=='en' else 'Kontaktas leidimams gauti: ')+F['contact'])
    for i,c in enumerate(C[lang],1):h(str(i)+' '+c['heading'],'clause_heading_'+str(i));p(c['body'],key='clause_body_'+str(i))
for lang in ['en','lt']:
    coverage[lang].append('provenance_title')
    page(T['provenance_title'][lang],('Version evidence  |  ' if lang=='en' else 'Versijų patikra  |  ')+F['date'])
    unit('provenance_intro');unit('repository',True)
    link(F['repo'],'https://github.com/'+F['repo']);p(F['branch'])
    for key in ['prior_commit','first_license_commit','runtime_commit']:
        unit(key,True);mono(F[key])
    unit('git_limit')
    unit('artifact_hashes',True)
    for a in F['artifacts']:
        pp=p(a['file']);pp.paragraph_format.keep_with_next=True
        for rr in pp.runs:rr.font.size=Pt(9)
        mono(a['sha256'])
    unit('artifact_limit');unit('declarations_title',True);unit('declarations')
    unit('third_party_title',True);unit('third_party')
    unit('sources_title',True)
    link('Open Source Initiative  MIT License','https://opensource.org/license/mit')
    link('WIPO  Copyright','https://www.wipo.int/en/web/copyright/protection')
    link('Git  verify commit','https://git-scm.com/docs/git-verify-commit')
for lang in ['en','lt']:
    coverage[lang].append('signature_title')
    page(T['signature_title'][lang],('Final distribution  |  Ed25519  |  1.1' if lang=='en' else 'Galutinis platinimo paketas  |  Ed25519  |  1.1'))
    unit('signature_intro');unit('fingerprint_title',True);mono(K['fingerprint_sha256'])
    for key,value in [('signer_id',K['signer_id']),('namespace',K['namespace']),('key_date',K['created_at_utc'])]:
        p(T[key][lang]+': '+value,key=key)
    unit('key_origin_title',True);unit('key_origin');unit('key_rotation');unit('trust')
    unit('verify_title',True);unit('verify_intro')
    for line in ['A='+F['archive'],'ssh-keygen -Y verify -f allowed_signers \\',
        '  -I '+K['signer_id']+' \\','  -n '+K['namespace']+' -s "$A.sig" < "$A"']:
        pp=mono(line);pp.paragraph_format.space_after=Pt(0)
    unit('verify_result');unit('receipt_title',True);unit('receipt')
    link('OpenSSH  ssh keygen','https://man.openbsd.org/ssh-keygen')
assert coverage['en']==coverage['lt']
for key in T:assert key in coverage['en'],key
doc=OUT/'Grandiniu_LD_licencija_EN_LT_ir_Git_patikra.docx';d.save(doc)
(R/'evidence/LANGUAGE_COVERAGE.json').write_text(json.dumps({'status':'PASS','equivalent_units_per_language':len(coverage['en']),'unit_ids':coverage['en'],'paired_clause_count':7,'same_fact_source':True},indent=2)+'\n')
print(json.dumps({'document':str(doc),'units_per_language':len(coverage['en'])}))
